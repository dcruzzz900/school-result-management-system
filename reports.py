"""Generic CSV/Excel writers used by the Reports & Analytics section.

Every report builds its data as a plain (headers, rows) pair, then hands it
to one of these two functions to produce the actual downloadable file. This
keeps the export format decoupled from how each report's data is computed.
"""
import io
import csv

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build_csv(headers, rows):
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    return io.BytesIO(buf.getvalue().encode("utf-8"))


def build_xlsx(title, headers, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] if title else "Report"  # Excel sheet names cap at 31 chars

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill

    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    # Auto-size columns roughly based on content length
    for col in range(1, len(headers) + 1):
        max_len = len(str(headers[col - 1]))
        for row in rows:
            if col - 1 < len(row) and row[col - 1] is not None:
                max_len = max(max_len, len(str(row[col - 1])))
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 3, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
